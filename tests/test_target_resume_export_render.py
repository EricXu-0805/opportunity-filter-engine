"""Real standard-v1 binary outputs: Unicode, pagination and editable full fonts."""
from __future__ import annotations

import io
import json
import re
import zipfile
from copy import deepcopy
from pathlib import Path
from uuid import UUID

import pytest
from lxml import etree

from backend.lib import target_resume_export as renderer
from backend.lib.target_resume_export_schema import ExportError

FIXTURE = Path(__file__).parent / 'fixtures' / 'target-resume-export-golden.json'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P = 'http://schemas.openxmlformats.org/package/2006/relationships'


def sample(text='Student 中文 😀'):
    return {'version': 1, 'template': 'standard-v1', 'locale': 'en', 'page_size': 'letter',
            'sections': [{'kind': 'basics', 'heading': '', 'blocks': [{'lines': [{'role': 'name', 'label': '', 'text': text}]}]}]}


def pdf_reader(data):
    from pypdf import PdfReader
    return PdfReader(io.BytesIO(data), strict=True)


def flat(text):
    return re.sub(r'\s', '', text)


@pytest.mark.parametrize('page_size,dimensions', [('letter', (612, 792)), ('a4', (595.28, 841.89))])
def test_pdf_true_unicode_text_and_page_dimensions(page_size, dimensions):
    value = sample('姓名 Student 中文 😀 <script>& text')
    value['page_size'] = page_size
    data = renderer.render_export(value, 'pdf')
    pdf = pdf_reader(data)
    assert flat('姓名 Student 中文 😀 <script>& text') in flat(pdf.pages[0].extract_text())
    box = pdf.pages[0].mediabox
    assert float(box.width) == pytest.approx(dimensions[0], abs=0.1)
    assert float(box.height) == pytest.approx(dimensions[1], abs=0.1)
    fonts = [f.get_object() for page in pdf.pages for f in page['/Resources']['/Font'].values()]
    assert any(b'D83DDE00' in font['/ToUnicode'].get_object().get_data() for font in fonts if '/ToUnicode' in font)
    assert any('/FontFile2' in font['/DescendantFonts'][0].get_object()['/FontDescriptor'] for font in fonts if '/DescendantFonts' in font)


def test_pdf_long_single_paragraph_and_unbroken_token_keep_tail_and_all_content():
    content = ('English text 中文 paragraph. ' * 1300) + ' LongTokenStart' + 'x' * 3000 + 'LongTokenEnd TAIL结束'
    value = sample('Student')
    value['sections'].append({'kind': 'activities', 'heading': '', 'blocks': [{'lines': [{'role': 'experience', 'label': '', 'text': content}]}]})
    data = renderer.render_export(value, 'pdf')
    pdf = pdf_reader(data)
    extracted = ''.join(page.extract_text() for page in pdf.pages)
    assert len(pdf.pages) > 3
    assert flat(content) in flat(extracted)
    assert 'TAIL结束' in flat(pdf.pages[-1].extract_text())
    assert all(page.extract_text().strip() for page in pdf.pages)


def test_pdf_signed_golden_preserves_selected_order_and_unicode():
    value = json.loads(FIXTURE.read_text())['projection']
    before = deepcopy(value)
    data = renderer.render_export(value, 'pdf')
    extracted = flat(''.join(page.extract_text() for page in pdf_reader(data).pages))
    cursor = 0
    for section in value['sections']:
        for block in section['blocks']:
            for line in block['lines']:
                needle = flat(line['text'])
                cursor = extracted.index(needle, cursor) + len(needle)
    assert value == before
    assert '徐同学😀' in extracted


def test_docx_full_fonts_relationships_and_editable_unicode():
    from docx import Document
    from fontTools.ttLib import TTFont
    value = sample('  姓名 Student 中文 😀 <script>& text\tline\r\nnext  ')
    value['sections'].append({'kind': 'activities', 'heading': '', 'blocks': [{'lines': [{'role': 'experience', 'label': '', 'text': 'A sentence that can be edited.'}]}]})
    data = renderer.render_export(value, 'docx')
    document = Document(io.BytesIO(data))
    assert '<script>& text' in document.paragraphs[0].text
    assert '\t' in document.paragraphs[0].text and '😀' in document.paragraphs[0].text
    assert document.paragraphs[0].text.count('\n') == 1  # CRLF is one visual break.
    assert document.paragraphs[0].text.startswith('  ') and document.paragraphs[0].text.endswith('  ')
    document.paragraphs[-1].add_run(' Edited after export 中文.')
    modified = io.BytesIO()
    document.save(modified)
    assert Document(io.BytesIO(modified.getvalue())).paragraphs[-1].text.endswith('Edited after export 中文.')
    z = zipfile.ZipFile(io.BytesIO(data))
    fonts = etree.fromstring(z.read('word/fontTable.xml'))
    relationships = etree.fromstring(z.read('word/_rels/fontTable.xml.rels'))
    by_id = {node.get('Id'): node for node in relationships}
    embedded = fonts.findall(f'{{{W}}}font/{{{W}}}embedRegular')
    assert len(embedded) == 2
    for node, asset in zip(embedded, renderer.fonts(), strict=True):
        assert node.get(f'{{{W}}}subsetted') == '0'
        rel = by_id[node.get(f'{{{R}}}id')]
        assert rel.get('Type') == R + '/font'
        assert rel.get('TargetMode') is None
        encrypted = bytearray(z.read('word/' + rel.get('Target')))
        key = UUID(node.get(f'{{{W}}}fontKey')).bytes[::-1]
        for i in range(32):
            encrypted[i] ^= key[i % 16]
        assert bytes(encrypted) == asset.data  # Whole programs, not today's text subset.
        with TTFont(io.BytesIO(encrypted)) as font:
            assert 'glyf' in font and 'fvar' not in font
            assert frozenset(font.getBestCmap()) == asset.codepoints
    settings = etree.fromstring(z.read('word/settings.xml'))
    assert settings.find(f'{{{W}}}embedTrueTypeFonts').get(f'{{{W}}}val') == 'true'
    assert settings.find(f'{{{W}}}saveSubsetFonts').get(f'{{{W}}}val') == 'false'
    assert not any(name.endswith(('vbaProject.bin', '.html')) for name in z.namelist())


@pytest.mark.parametrize('output', ['pdf', 'docx'])
@pytest.mark.parametrize('text', ['\U0010ffff', '👩\u200d💻', '👍🏽', '🇺🇸', '☀\ufe0f'])
def test_unsupported_glyph_or_sequence_is_rejected_not_dropped(output, text):
    with pytest.raises(ExportError, match='^unsupported_glyph$'):
        renderer.render_export(sample(text), output)


def test_missing_font_assets_are_safe(monkeypatch, tmp_path):
    renderer.fonts.cache_clear()
    monkeypatch.setattr(renderer, 'FONT_DIRECTORY', tmp_path)
    try:
        with pytest.raises(ExportError, match='^fonts_unavailable$'):
            renderer.render_export(sample(), 'pdf')
    finally:
        renderer.fonts.cache_clear()


def test_renderer_uses_text_not_markup_or_external_fetch(monkeypatch):
    import urllib.request
    monkeypatch.setattr(urllib.request, 'urlopen', lambda *_a, **_kw: pytest.fail('export attempted a fetch'))
    value = sample('<img src="https://example.invalid/private"> & entity')
    value['sections'][0]['blocks'][0]['lines'].append({'role': 'url', 'label': '', 'text': 'javascript:alert(1)'})
    for output in ('pdf', 'docx'):
        data = renderer.render_export(value, output)
        if output == 'pdf':
            pdf = pdf_reader(data)
            assert '<imgsrc="https://example.invalid/private">&entity' in flat(pdf.pages[0].extract_text())
            assert not pdf.pages[0].get('/Annots')
        else:
            z = zipfile.ZipFile(io.BytesIO(data))
            xml = z.read('word/document.xml')
            assert b'&lt;img' in xml and b'<img ' not in xml
            assert b'TargetMode="External"' not in z.read('word/_rels/document.xml.rels')


def test_docx_and_pdf_links_preserve_safe_target_as_annotations_only():
    value = sample('Student')
    url = 'https://example.edu/path?q=a&mode=b'
    value['sections'][0]['blocks'][0]['lines'].append({'role': 'url', 'label': 'Portfolio', 'text': url})
    pdf = pdf_reader(renderer.render_export(value, 'pdf'))
    assert any(annotation.get_object()['/A']['/URI'] == url for annotation in pdf.pages[0]['/Annots'])
    z = zipfile.ZipFile(io.BytesIO(renderer.render_export(value, 'docx')))
    relationships = etree.fromstring(z.read('word/_rels/document.xml.rels'))
    assert any(node.get('Target') == url and node.get('TargetMode') == 'External' for node in relationships)
    assert url in ''.join(etree.fromstring(z.read('word/document.xml')).itertext())


def test_output_size_guards_do_not_trim_or_return_partial_files(monkeypatch):
    buffer = renderer.CappedBuffer()
    buffer.seek(renderer.MAX_FILE_BYTES)
    with pytest.raises(ExportError, match='^export_too_large$'):
        buffer.write(b'x')
    monkeypatch.setattr(renderer, 'MAX_FILE_BYTES', 1024)
    with pytest.raises(ExportError, match='^export_too_large$'):
        renderer.render_export(sample(), 'pdf')


def test_pdf_page_limit_refuses_next_page_without_returning_a_partial_file(monkeypatch):
    from fpdf import FPDF
    monkeypatch.setattr(renderer, 'MAX_PDF_PAGES', 2)
    created = []
    original = FPDF.add_page

    def record(pdf, *args, **kwargs):
        result = original(pdf, *args, **kwargs)
        created.append(pdf.page_no())
        return result

    monkeypatch.setattr(FPDF, 'add_page', record)
    value = sample('x\n' * 150)
    before = deepcopy(value)
    with pytest.raises(ExportError, match='^export_too_large$'):
        renderer.render_export(value, 'pdf')
    assert created == [1, 2]  # Page 3 is rejected before FPDF creates it.
    assert value == before


@pytest.mark.parametrize('output', ['pdf', 'docx'])
def test_expired_worker_deadline_stops_before_loading_fonts(monkeypatch, output):
    monkeypatch.setattr(renderer, 'monotonic', lambda: 2.0)
    monkeypatch.setattr(renderer, 'fonts', lambda: pytest.fail('expired work started font loading'))
    with pytest.raises(ExportError, match='^export_timeout$'):
        renderer.render_export(sample(), output, deadline=1.0)


def test_glyph_preflight_checks_deadline_after_bounded_shaping_chunk(monkeypatch):
    import uharfbuzz as hb
    clock, lengths = [0.0], []
    monkeypatch.setattr(renderer, 'monotonic', lambda: clock[0])
    original = hb.shape

    def shape(font, buffer):
        lengths.append(len(buffer.glyph_infos))
        original(font, buffer)
        clock[0] = 2.0

    monkeypatch.setattr(hb, 'shape', shape)
    with pytest.raises(ExportError, match='^export_timeout$'):
        renderer.render_export(sample('x' * 12000), 'pdf', deadline=1.0)
    assert lengths == [renderer.GLYPH_CHECK_CHARS]


def test_pdf_automatic_page_break_checks_deadline_inside_single_paragraph(monkeypatch):
    from fpdf import FPDF
    clock, pages = [0.0], []
    monkeypatch.setattr(renderer, 'monotonic', lambda: clock[0])
    original = FPDF.add_page

    def add_page(pdf, *args, **kwargs):
        result = original(pdf, *args, **kwargs)
        pages.append(pdf.page_no())
        if len(pages) == 2:
            clock[0] = 2.0
        return result

    monkeypatch.setattr(FPDF, 'add_page', add_page)
    with pytest.raises(ExportError, match='^export_timeout$'):
        renderer.render_export(sample('x\n' * 150), 'pdf', deadline=1.0)
    assert pages == [1, 2]


def test_docx_iteration_and_zip_writes_check_deadline(monkeypatch):
    clock = [0.0]
    monkeypatch.setattr(renderer, 'monotonic', lambda: clock[0])
    original = renderer.font_runs
    seen = []

    def runs(text, assets, deadline=None):
        for index, run in original(text, assets, deadline):
            seen.append(run)
            yield index, run
            clock[0] = 2.0

    monkeypatch.setattr(renderer, 'font_runs', runs)
    with pytest.raises(ExportError, match='^export_timeout$'):
        renderer.render_docx(sample('English😀'), renderer.fonts(), deadline=1.0)
    assert seen == ['English', '😀']
    buffer = renderer.CappedBuffer(deadline=1.0)
    with pytest.raises(ExportError, match='^export_timeout$'):
        buffer.write(b'PK')
    assert buffer.getvalue() == b''


@pytest.mark.parametrize('position', ['heading', 'body'])
def test_wrapped_pdf_first_line_keeps_natural_word_spacing(position):
    value = sample('Student')
    if position == 'heading':
        value['sections'][0]['heading'] = 'Research notes ' + 'x' * 100
    else:
        value['sections'][0]['blocks'][0]['lines'] = [
            {'role': 'other', 'label': 'Additional notes 补充说明', 'text': 'x' * 120},
        ]
    page = pdf_reader(renderer.render_export(value, 'pdf')).pages[0]
    # Read actual text placement. The long following token wraps, leaving a
    # short first row; justification formerly stretched its words across 500pt.
    positions = []
    for operands, operator in page.get_contents().operations:
        if operator == b'Td' and not positions:
            positions.append(float(operands[0]))
        elif operator == b'Tm' and positions:
            positions.append(float(operands[4]))
        elif operator == b'ET' and positions:
            break
    assert len(positions) >= 2
    assert max(positions) - positions[0] < 160, positions
