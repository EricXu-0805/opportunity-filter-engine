"""Text-only standard-v1 PDF/DOCX renderers using pinned local font assets.

No user markup, remote resources, source material, or current profile is read.
Both outputs keep the submitted section/block/line order and current wording.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from time import monotonic
from urllib.parse import quote, urlsplit
from uuid import uuid4

from backend.lib.target_resume_export_schema import MAX_FILE_BYTES, ExportError

FONT_DIRECTORY = Path(__file__).resolve().parents[1] / 'assets' / 'resume-fonts'
FONT_FILES = ('NotoSansCJKsc-Regular.ttf', 'NotoEmoji-Regular.ttf')
MAX_PDF_PAGES = 100
GLYPH_CHECK_CHARS = 4096


def check_deadline(deadline):
    if deadline is not None and monotonic() >= deadline:
        raise ExportError('export_timeout')

HEADINGS = {
    'en': {'basics': '', 'education': 'Education', 'activities': 'Experience', 'publications': 'Publications', 'skills': 'Skills', 'other': 'Additional information'},
    'zh': {'basics': '', 'education': '教育经历', 'activities': '项目与经历', 'publications': '论文与发表', 'skills': '技能', 'other': '其他信息'},
}
LABELS = {
    'en': {'email': 'Email', 'phone': 'Phone', 'location': 'Location', 'school': 'School', 'degree': 'Degree', 'field': 'Field',
           'start': 'Start', 'end': 'End', 'organization': 'Organization', 'authors': 'Authors', 'venue': 'Venue', 'date': 'Date',
           'publication_status': 'Status', 'url': 'Link', 'doi': 'DOI'},
    'zh': {'email': '邮箱', 'phone': '电话', 'location': '地点', 'school': '学校', 'degree': '学位', 'field': '专业',
           'start': '开始', 'end': '结束', 'organization': '机构', 'authors': '作者', 'venue': '刊物或会议', 'date': '日期',
           'publication_status': '状态', 'url': '链接', 'doi': 'DOI'},
}


@dataclass(frozen=True)
class FontAsset:
    path: Path
    name: str
    data: bytes
    codepoints: frozenset[int]


@lru_cache(maxsize=1)
def fonts() -> tuple[FontAsset, FontAsset]:
    try:
        from fontTools.ttLib import TTFont
        result = []
        for filename in FONT_FILES:
            path = FONT_DIRECTORY / filename
            data = path.read_bytes()
            with TTFont(io.BytesIO(data)) as font:
                # The same complete static glyf fonts are embedded in editable
                # DOCX. Preview/print-only and restricted fonts are unsuitable.
                if 'glyf' not in font or 'fvar' in font or (font['OS/2'].fsType & 0x0006):
                    raise ExportError('fonts_unavailable')
                name = font['name'].getDebugName(1)
                if not name or not font.getBestCmap():
                    raise ExportError('fonts_unavailable')
                result.append(FontAsset(path, name, data, frozenset(font.getBestCmap())))
        return tuple(result)
    except ExportError:
        raise
    except Exception:
        raise ExportError('fonts_unavailable') from None


def font_runs(text: str, assets: tuple[FontAsset, FontAsset], deadline=None):
    """Fallback by supported glyph; never silently output .notdef boxes.

    standard-v1 supports individual monochrome emoji. Compound emoji/IVS
    sequences need an independently verified shaping-and-extraction contract;
    reject them intact until then, rather than deleting joiners/modifiers.
    """
    current, buffer = None, []
    for offset, char in enumerate(text):
        if offset % 1024 == 0:
            check_deadline(deadline)
        point = ord(char)
        if (point in (0x200D, 0x20E3, 0xFE0E, 0xFE0F) or 0x1F3FB <= point <= 0x1F3FF
                or 0x1F1E6 <= point <= 0x1F1FF or 0xE0020 <= point <= 0xE007F or 0xE0100 <= point <= 0xE01EF):
            raise ExportError('unsupported_glyph')
        index = 0 if char in '\r\n\t' or point in assets[0].codepoints else 1 if point in assets[1].codepoints else None
        if index is None:
            raise ExportError('unsupported_glyph')
        if current is not None and index != current:
            yield current, ''.join(buffer)
            buffer = []
        current = index
        buffer.append(char)
    if buffer:
        yield current, ''.join(buffer)


def heading(section, locale):
    return section['heading'] or HEADINGS[locale][section['kind']]


def line_text(line, locale):
    label = line['label'] or LABELS[locale].get(line['role'], '')
    return f'{label}: {line["text"]}' if label else line['text']


def validate_glyphs(projection, assets, deadline=None):
    import uharfbuzz as hb
    hb_fonts = [hb.Font(hb.Face(asset.data)) for asset in assets]
    for section in projection['sections']:
        check_deadline(deadline)
        texts = [heading(section, projection['locale'])]
        texts.extend(line_text(line, projection['locale']) for block in section['blocks'] for line in block['lines'])
        for text in texts:
            check_deadline(deadline)
            for index, run in font_runs(text, assets, deadline):
                for part in re.split(r'[\r\n\t]', run):
                    # Bounded preflight chunks only. The actual renderer still
                    # receives the full original paragraph, without truncation.
                    for start in range(0, len(part), GLYPH_CHECK_CHARS):
                        check_deadline(deadline)
                        buffer = hb.Buffer()
                        buffer.add_str(part[start:start + GLYPH_CHECK_CHARS])
                        buffer.guess_segment_properties()
                        hb.shape(hb_fonts[index], buffer)
                        check_deadline(deadline)
                        if any(info.codepoint == 0 for info in buffer.glyph_infos):
                            raise ExportError('unsupported_glyph')


def safe_link(line):
    """Only explicit ordinary links become annotations; never fetch a URL."""
    value = line['text']
    if not value or any(char.isspace() or ord(char) < 0x20 for char in value):
        return None
    if line['role'] == 'email' and re.fullmatch(r'[A-Za-z0-9.!#$%&\'*+/=?^_`{|}~-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}', value):
        return 'mailto:' + quote(value, safe='@.+-_')
    if line['role'] == 'doi' and re.fullmatch(r'10\.\d{4,9}/[^\s]+', value):
        return 'https://doi.org/' + quote(value, safe='/.:()-_')
    if line['role'] not in ('url', 'doi'):
        return None
    try:
        parsed = urlsplit(value)
        if parsed.scheme not in ('https', 'http') or not parsed.hostname or parsed.username or parsed.password:
            return None
        _port = parsed.port  # Reject malformed port values without exposing them.
        if any(char in value for char in '<>"\\'):
            return None
        return value
    except ValueError:
        return None


def render_pdf(projection, assets, deadline=None):
    from fpdf import FPDF
    from fpdf.enums import MethodReturnValue, XPos, YPos

    class BoundedPDF(FPDF):
        def add_page(self, *args, **kwargs):
            check_deadline(deadline)
            if self.page_no() >= MAX_PDF_PAGES:
                raise ExportError('export_too_large')
            return super().add_page(*args, **kwargs)

    check_deadline(deadline)
    pdf = BoundedPDF(format='Letter' if projection['page_size'] == 'letter' else 'A4')
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(True, margin=16)
    pdf.set_title('Resume')
    pdf.set_author('')
    pdf.add_font('ResumeSans', fname=str(assets[0].path))
    check_deadline(deadline)
    pdf.add_font('ResumeEmoji', fname=str(assets[1].path))
    check_deadline(deadline)
    pdf.set_fallback_fonts(['ResumeEmoji'], exact_match=False)
    pdf.set_text_shaping(True)
    pdf.add_page()
    pdf.set_text_color(20, 28, 39)
    for section_index, section in enumerate(projection['sections']):
        check_deadline(deadline)
        title = heading(section, projection['locale'])
        if title:
            pdf.set_font('ResumeSans', size=12)
            title_text = title.replace('\r\n', '\n').replace('\r', '\n').replace('\t', '    ')
            height = pdf.multi_cell(0, 6, title_text, align='L', dry_run=True, output=MethodReturnValue.HEIGHT)
            if pdf.will_page_break(height + 6):
                pdf.add_page()
            elif section_index:
                pdf.ln(3)
            pdf.multi_cell(0, 6, title_text, align='L', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(1)
        for block in section['blocks']:
            for line in block['lines']:
                check_deadline(deadline)
                text = line_text(line, projection['locale'])
                # CRLF is one visual break; native PDF text has no tab control.
                # This is layout only: the signed projection is never changed.
                text = text.replace('\r\n', '\n').replace('\r', '\n').replace('\t', '    ')
                size, leading = (18, 8) if line['role'] == 'name' else (10.5, 5.3)
                pdf.set_font('ResumeSans', size=size)
                pdf.multi_cell(0, leading, text, align='L', new_x=XPos.LMARGIN, new_y=YPos.NEXT, link=safe_link(line) or '')
                check_deadline(deadline)
                pdf.ln(0.6)
            pdf.ln(2)
    check_deadline(deadline)
    data = bytes(pdf.output())
    check_deadline(deadline)
    if len(data) > MAX_FILE_BYTES:
        raise ExportError('export_too_large')
    return data


class CappedBuffer(io.BytesIO):
    def __init__(self, deadline=None):
        super().__init__()
        self.deadline = deadline

    def write(self, data):
        check_deadline(self.deadline)
        if self.tell() + len(data) > MAX_FILE_BYTES:
            raise ExportError('export_too_large')
        return super().write(data)


def embed_docx_fonts(document, assets, deadline=None):
    """Embed full font programs, not subsets limited to today's résumé text."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table = document.part.part_related_by(RT.FONT_TABLE)
    # fontTable is a generic OPC Part in python-docx; do not depend on a
    # private loaded XML object that may not exist across versions.
    from lxml import etree
    root = etree.fromstring(table.blob, parser=etree.XMLParser(resolve_entities=False, no_network=True))
    for index, asset in enumerate(assets):
        check_deadline(deadline)
        key = uuid4()
        font_bytes = bytearray(asset.data)
        mask = key.bytes[::-1]
        for offset in range(32):
            font_bytes[offset] ^= mask[offset % 16]
        part = Part(PackURI(f'/word/fonts/resume-{index}.odttf'),
                    'application/vnd.openxmlformats-officedocument.obfuscatedFont', bytes(font_bytes), document.part.package)
        rid = table.relate_to(part, RT.FONT)
        matches = root.findall(qn('w:font'))
        for old in matches:
            if old.get(qn('w:name')) == asset.name:
                root.remove(old)
        font = OxmlElement('w:font')
        font.set(qn('w:name'), asset.name)
        embedded = OxmlElement('w:embedRegular')
        embedded.set(qn('r:id'), rid)
        embedded.set(qn('w:fontKey'), '{' + str(key).upper() + '}')
        embedded.set(qn('w:subsetted'), '0')
        font.append(embedded)
        root.append(font)
    table._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    settings = document.settings.element
    for tag, value in [('embedTrueTypeFonts', 'true'), ('saveSubsetFonts', 'false')]:
        node = settings.find(qn('w:' + tag))
        if node is None:
            node = OxmlElement('w:' + tag)
            settings.append(node)
        node.set(qn('w:val'), value)


def render_docx(projection, assets, deadline=None):
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor

    check_deadline(deadline)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = (Inches(8.5), Inches(11)) if projection['page_size'] == 'letter' else (Mm(210), Mm(297))
    section.top_margin = section.bottom_margin = Mm(16)
    section.left_margin = section.right_margin = Mm(18)
    normal = document.styles['Normal']
    normal.font.name = assets[0].name
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(20, 28, 39)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), assets[0].name)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.widow_control = True
    for key in ('author', 'last_modified_by', 'subject', 'comments', 'keywords', 'category', 'description'):
        if hasattr(document.core_properties, key):
            setattr(document.core_properties, key, '')
    document.core_properties.title = 'Resume'

    def add_text(paragraph, text, size=10.5, link=None):
        parent = paragraph._p
        if link:
            hyperlink = OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'), document.part.relate_to(link, RT.HYPERLINK, is_external=True))
            parent.append(hyperlink)
            parent = hyperlink
        for index, value in font_runs(text, assets, deadline):
            check_deadline(deadline)
            # python-docx maps CR and LF separately; treat CRLF as one
            # visual break while leaving the signed input untouched.
            run = paragraph.add_run(value.replace('\r\n', '\n').replace('\r', '\n'))
            run.font.name = assets[index].name
            run.font.size = Pt(size)
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), assets[index].name)
            if link:
                parent.append(run._r)
    for item in projection['sections']:
        check_deadline(deadline)
        title = heading(item, projection['locale'])
        if title:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.keep_with_next = True
            add_text(paragraph, title, 12)
        for block in item['blocks']:
            for index, line in enumerate(block['lines']):
                check_deadline(deadline)
                paragraph = document.add_paragraph()
                # Whole long paragraphs must remain splittable over pages.
                paragraph.paragraph_format.keep_together = False
                paragraph.paragraph_format.keep_with_next = False
                paragraph.paragraph_format.space_after = Pt(6 if index == len(block['lines']) - 1 else 3)
                add_text(paragraph, line_text(line, projection['locale']), 18 if line['role'] == 'name' else 10.5, safe_link(line))
    embed_docx_fonts(document, assets, deadline)
    output = CappedBuffer(deadline)
    document.save(output)
    check_deadline(deadline)
    return output.getvalue()


def render_export(projection, output_format, *, deadline=None):
    try:
        check_deadline(deadline)
        assets = fonts()
        check_deadline(deadline)
        validate_glyphs(projection, assets, deadline)
        return render_pdf(projection, assets, deadline) if output_format == 'pdf' else render_docx(projection, assets, deadline)
    except ExportError:
        raise
    except (ImportError, FileNotFoundError):
        raise ExportError('fonts_unavailable') from None
    except Exception:
        raise ExportError('export_failed') from None
